"""Extraction, chunking and planning tests. The samples come from scripts/make_samples.py."""
import datetime as dt
import io
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook
from pptx import Presentation

import deckray
from deckray import ingest, pptx
from deckray.chunks import select_context, sources
from deckray.llm import validate_plan

SAMPLES = Path(__file__).resolve().parent.parent / "samples"
DECK = ingest("deck.pptx", (SAMPLES / "Solar Program Q3 Review.pptx").read_bytes())
BOOK = ingest("book.xlsx", (SAMPLES / "Solar Sites Directory.xlsx").read_bytes())
slides = {s["index"]: s for s in DECK["extraction"]["slides"]}
sheets = {s["name"]: s for s in BOOK["extraction"]["sheets"]}


def test_slides_follow_presentation_order_past_slide_9():
    prs = Presentation()
    for i in range(12):
        prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = f"Slide {i + 1}"
    buf = io.BytesIO(); prs.save(buf)
    titles = [s["title"] for s in pptx.extract(zipfile.ZipFile(buf))["slides"]]
    assert titles == [f"Slide {i + 1}" for i in range(12)]  # filename sort would put slide10 after slide1


def test_list_hierarchy_and_inherited_bullets():
    assert slides[2]["lines"][:4] == ["1. Installations by region", "2. Generation results",
                                      "  1. Rooftop vs carport output", "  2. Weather impact in August"]
    assert slides[5]["lines"][:2] == ["• Inverter supply delays", "  • Mitigation: dual-source from two vendors"]


def test_table_merged_cells_fill_only_merged_areas():
    rows = slides[3]["tables"][0]["rows"]
    assert [r["Region"] for r in rows] == ["Spain", "Spain", "Portugal", "Portugal", "India"]


def test_chart_keeps_every_point_of_every_series():
    chart = slides[4]["charts"][0]
    assert chart["title"] == "Monthly generation, Q3 2026"
    assert {s["name"]: [p["value"] for p in s["points"]] for s in chart["series"]} == {
        "Rooftop (MWh)": [612, 548, 590], "Carport (MWh)": [201, 176, 214]}


def test_notes_links_and_image_alt_text():
    assert slides[5]["notes"].startswith("Stress that the inverter delay")
    assert slides[5]["links"] == [{"text": "Full risk register", "url": "https://example.com/brightfield/risk-register"}]
    assert slides[6]["images"][0]["alt"].startswith("Map of active sites") and slides[6]["images"][0]["part"]


def test_smartart_tree_comes_from_connections():
    ns = 'xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    pts = "".join(f'<dgm:pt modelId="{m}"><dgm:t><a:p><a:r><a:t>{t}</a:t></a:r></a:p></dgm:t></dgm:pt>'
                  for m, t in (("1", "Plan"), ("2", "Build"), ("3", "Design"), ("4", "Review")))
    cxns = '<dgm:cxn srcId="0" destId="1"/><dgm:cxn srcId="0" destId="2"/><dgm:cxn srcId="1" destId="3"/><dgm:cxn srcId="1" destId="4"/>'
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("data.xml", f'<dgm:dataModel {ns}><dgm:ptLst><dgm:pt modelId="0" type="doc"/>{pts}</dgm:ptLst><dgm:cxnLst>{cxns}</dgm:cxnLst></dgm:dataModel>')
    tree = pptx._smartart(zipfile.ZipFile(buf), "data.xml")
    assert [(n["text"], [c["text"] for c in n["children"]]) for n in tree] == [("Plan", ["Design", "Review"]), ("Build", [])]


def test_excel_header_found_under_title_rows_despite_optional_column():
    s = sheets["Sites"]
    assert s["header_row"] == 4 and s["headers"][:3] == ["Site ID", "Country", "City"]
    assert s["description"].startswith("Brightfield Energy (fictional)")


def test_blank_header_keeps_column_positions():
    row = next(r for r in sheets["Sites"]["rows"] if r["fields"]["Site ID"] == "ES-02")
    assert row["fields"]["Column G"] == "Roof reinforced" and row["fields"]["Status"] == "Online"


def test_merged_cells_unmerged_in_data_rows_and_dates_links_hidden():
    rows = sheets["Sites"]["rows"]
    assert [r["fields"]["Country"] for r in rows].count("Spain") == 5
    assert rows[0]["fields"]["Commissioned"] == "2025-03-12"
    assert rows[0]["links"] == {"Site ID": "https://example.com/brightfield/monitor/ES-01"}
    assert sheets["Assumptions"]["hidden"] and not sheets["Vendors"]["hidden"]


def test_naive_baseline_loses_what_structure_keeps():
    naive = " ".join(c["text"] for c in BOOK["naive"])
    assert "ES-02 Seville" in naive  # the country is gone: only the merge's first cell had it
    assert any("ES-02" in c["text"] and "Country: Spain" in c["text"] for c in BOOK["chunks"])


def test_ingest_rejects_bad_input(monkeypatch):
    with pytest.raises(deckray.IngestError, match="only .pptx"):
        ingest("notes.txt", b"hi")
    with pytest.raises(deckray.IngestError, match="isn't a valid Office file"):
        ingest("old.xlsx", b"not a zip")
    monkeypatch.setattr(deckray, "MAX_UNZIPPED", 1000)
    with pytest.raises(deckray.IngestError, match="expands to too much"):
        ingest("book.xlsx", (SAMPLES / "Solar Sites Directory.xlsx").read_bytes())


def test_plan_validation_drops_invented_sheets_columns_and_values():
    src = sources({"book.xlsx": BOOK["extraction"]})
    plan = validate_plan({"route": "xlsx", "sheet": "sites", "filters": {"country": "spain", "Status": "Broken", "Owner": "x"}}, src)
    assert plan["sheet"] == "Sites" and plan["filters"] == {"Country": "Spain"}
    assert validate_plan({"route": "nonsense", "sheet": "Nope"}, src) == {"route": "both", "sheet": None, "filters": {}, "query": ""}


def test_select_context_applies_filters_and_pulls_slide_siblings():
    chunks = DECK["chunks"] + BOOK["chunks"]
    scores = [1.0 if c["id"].endswith("s5-text") else 0.0 for c in chunks]
    ctx = select_context(chunks, scores, {"route": "both"}, k=1)
    assert {c["id"] for c in ctx} == {c["id"] for c in DECK["chunks"] if c["where"].get("slide") == 5}
    rows = select_context(chunks, [0.0] * len(chunks), {"route": "xlsx", "sheet": "Sites", "filters": {"Country": "Spain", "Status": "Offline"}})
    assert [c["fields"]["Site ID"] for c in rows] == ["ES-03"]


async def _empty_post(url, headers, body):
    return {"choices": [{"message": {"content": ""}}]}


def test_empty_model_responses_are_errors_not_blank_answers():
    import asyncio
    from deckray import llm
    cfg = {"provider": "ollama", "model": "m"}
    plan = asyncio.run(llm.plan(cfg, _empty_post, "q", []))
    assert plan["source"] == "fallback" and "empty response" in plan["error"]
    with pytest.raises(RuntimeError, match="empty response"):
        asyncio.run(llm.answer(cfg, _empty_post, "q", [{"text": "x"}]))

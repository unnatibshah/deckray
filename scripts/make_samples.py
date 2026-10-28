"""Builds the fictional sample files used by the demo and the tests (dev only: needs python-pptx and pillow).

Run: python scripts/make_samples.py   -> samples/Solar Program Q3 Review.pptx, samples/Solar Sites Directory.xlsx
"""
import datetime as dt
import io
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent.parent / "samples"
OUT.mkdir(exist_ok=True)


def bullets(slide, items):
    """items: (level, text); the body placeholder supplies bullets, as in real decks."""
    tf = slide.placeholders[1].text_frame
    tf.text = items[0][1]
    for lvl, text in items[1:]:
        p = tf.add_paragraph()
        p.text, p.level = text, lvl


def deck():
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Rooftop Solar Program: Q3 2026 Review"
    s.placeholders[1].text = "Brightfield Energy (fictional) · Program office"

    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Agenda"
    tf = s.placeholders[1].text_frame
    for i, (lvl, text) in enumerate([(0, "Installations by region"), (0, "Generation results"),
                                     (1, "Rooftop vs carport output"), (1, "Weather impact in August"),
                                     (0, "Risks and mitigations"), (0, "Q4 priorities")]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text, p.level = text, lvl
        pPr = p._p.get_or_add_pPr()  # numbered list
        pPr.insert(0, pPr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}buAutoNum", {"type": "arabicPeriod"}))

    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Installations by region"
    rows = [("Region", "Quarter", "Sites", "Capacity (MW)"), ("Spain", "Q2", "14", "6.1"), ("", "Q3", "19", "8.4"),
            ("Portugal", "Q2", "6", "2.2"), ("", "Q3", "9", "3.5"), ("India", "Q3", "11", "5.0")]
    t = s.shapes.add_table(len(rows), 4, Inches(0.8), Inches(1.6), Inches(8.4), Inches(3)).table
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    t.cell(1, 0).merge(t.cell(2, 0))  # Spain spans Q2 and Q3
    t.cell(3, 0).merge(t.cell(4, 0))  # Portugal spans Q2 and Q3

    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Monthly generation"
    cd = CategoryChartData()
    cd.categories = ["July", "August", "September"]
    cd.add_series("Rooftop (MWh)", (612, 548, 590))
    cd.add_series("Carport (MWh)", (201, 176, 214))
    ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.8), Inches(1.5), Inches(8.4), Inches(4.5), cd).chart
    ch.has_title, ch.chart_title.text_frame.text = True, "Monthly generation, Q3 2026"

    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Risks and mitigations"
    bullets(s, [(0, "Inverter supply delays"), (1, "Mitigation: dual-source from two vendors"), (1, "Owner: procurement, due 15 Nov"),
                (0, "Grid connection backlog in Portugal"), (1, "Mitigation: pre-file connection requests 90 days earlier"),
                (0, "Panel soiling in dry months"), (1, "Mitigation: quarterly cleaning at India sites")])
    tb = s.shapes.add_textbox(Inches(0.8), Inches(6.6), Inches(8), Inches(0.5)).text_frame
    run = tb.paragraphs[0].add_run()
    run.text, run.font.size = "Full risk register", Pt(12)
    run.hyperlink.address = "https://example.com/brightfield/risk-register"
    s.notes_slide.notes_text_frame.text = ("Stress that the inverter delay is the only risk that threatens the Q4 target. "
                                          "Portugal's backlog is regulatory, not technical.")

    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Where we operate"
    img = Image.new("RGB", (640, 360), "#eef3f7")
    d = ImageDraw.Draw(img)
    for x, y in ((180, 150), (210, 170), (470, 220)):
        d.ellipse((x - 9, y - 9, x + 9, y + 9), fill="#f2a12b")
    buf = io.BytesIO(); img.save(buf, "PNG"); buf.seek(0)
    pic = s.shapes.add_picture(buf, Inches(1.5), Inches(1.6), Inches(7))
    pic._element.nvPicPr.cNvPr.set("descr", "Map of active sites: clusters in Spain and Portugal, one cluster in southern India")

    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Q4 priorities"
    bullets(s, [(0, "Commission 12 more sites in Spain"), (0, "Close the Portugal connection backlog"),
                (0, "Pilot battery storage at two carport sites"), (1, "Target: 15% evening self-consumption")])
    prs.save(OUT / "Solar Program Q3 Review.pptx")


def workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "Sites"
    ws["A1"] = "Brightfield Energy (fictional): rooftop solar sites"
    ws["A2"] = "As of 30 Sep 2026 · capacities in kW"
    ws.merge_cells("A1:H1")
    ws["A1"].font = Font(bold=True, size=14)
    headers = ["Site ID", "Country", "City", "Capacity (kW)", "Commissioned", "Panel vendor", None, "Status"]
    ws.append([])
    ws.append(headers)  # row 4; column G has no header but holds notes
    sites = [
        ("ES-01", "Spain", "Seville", 480, dt.date(2025, 3, 12), "Helio", "", "Online"),
        ("ES-02", "Spain", "Seville", 350, dt.date(2025, 6, 2), "Helio", "Roof reinforced", "Online"),
        ("ES-03", "Spain", "Valencia", 620, dt.date(2026, 1, 20), "SunPeak", "", "Offline"),
        ("ES-04", "Spain", "Madrid", 410, dt.date(2026, 7, 8), "SunPeak", "Inverter replaced Aug", "Online"),
        ("ES-05", "Spain", "Madrid", 300, dt.date(2026, 9, 1), "Helio", "", "Commissioning"),
        ("PT-01", "Portugal", "Lisbon", 260, dt.date(2025, 11, 4), "Lumen", "", "Online"),
        ("PT-02", "Portugal", "Porto", 310, None, "Lumen", "Waiting for grid connection", "Pending grid"),
        ("PT-03", "Portugal", "Porto", 280, None, "SunPeak", "Waiting for grid connection", "Pending grid"),
        ("IN-01", "India", "Chennai", 900, dt.date(2025, 9, 15), "Helio", "Quarterly cleaning", "Online"),
        ("IN-02", "India", "Chennai", 750, dt.date(2026, 2, 27), "Lumen", "", "Offline"),
        ("IN-03", "India", "Bengaluru", 540, dt.date(2026, 5, 19), "SunPeak", "", "Online"),
    ]
    for s in sites:
        ws.append(list(s))
    ws.merge_cells("B5:B9")    # Spain spans its five sites
    ws.merge_cells("B10:B12")  # Portugal
    ws.merge_cells("B13:B15")  # India
    for r in range(5, 16):
        ws.cell(r, 2).alignment = Alignment(vertical="center")
        ws.cell(r, 1).hyperlink = f"https://example.com/brightfield/monitor/{ws.cell(r, 1).value}"
    for c in "ABCDEFGH":
        ws.column_dimensions[c].width = 16

    v = wb.create_sheet("Vendors")
    v.append(["Vendor", "Panel model", "Efficiency (%)", "Warranty (years)", "Support contact"])
    for row in (("Helio", "H-440 Mono", 21.8, 25, "support@helio.example"), ("SunPeak", "SP Bifacial 455", 22.6, 30, "care@sunpeak.example"),
                ("Lumen", "L-410", 20.9, 20, "help@lumen.example")):
        v.append(list(row))

    a = wb.create_sheet("Assumptions")
    a.append(["Assumption", "Value", "Unit"])
    for row in (("Annual degradation", 0.5, "%"), ("Performance ratio", 0.81, "ratio"), ("Grid tariff", 0.11, "EUR/kWh")):
        a.append(list(row))
    a.sheet_state = "hidden"
    wb.save(OUT / "Solar Sites Directory.xlsx")


if __name__ == "__main__":
    deck()
    workbook()
    print("wrote", *sorted(p.name for p in OUT.iterdir()))

"""Excel extraction: header detection under title rows, merged cells unmerged everywhere, one record per row.

Columns keep their real positions (a blank header never shifts later columns), hidden sheets are flagged,
dates come out as ISO strings, and cell hyperlinks are kept.
"""
from __future__ import annotations

import datetime as dt
import re
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

MAX_ROWS, MAX_COLS = 5000, 60  # ponytail: demo-sized caps; stream with read_only mode for bigger sheets
EMPTY = {"", "nan", "n/a", "none", "-"}


def _clean(v: Any) -> str | None:
    if v is None:
        return None
    if isinstance(v, dt.datetime):
        v = v.date() if v.time() == dt.time() else v
    if isinstance(v, (dt.date, dt.datetime)):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return None if s.lower() in EMPTY else s


def key(header: str) -> str:
    return re.sub(r"_+", "_", re.sub(r"[^a-z0-9]", "_", header.lower())).strip("_") or "column"


LABEL = re.compile(r"^(?!-?[\d.,]+%?$)(?!\d{4}-\d{2}-\d{2})")  # not a number, percentage or ISO date


def _header_row(grid: list[list[str | None]], look: int = 10) -> int:
    """Header = first row of >=2 distinct text labels (no numbers or dates) with data below it.

    Falls back to "first row with >=2 distinct values that the next rows don't outgrow", which breaks when an
    optional column is filled in some data rows but not in the header row.
    """
    width = lambda r: sum(1 for v in r if v)
    for i in range(min(look, len(grid) - 1)):
        vals = [v for v in grid[i] if v]
        if len(vals) >= 2 and len(set(vals)) == len(vals) and all(LABEL.match(v) for v in vals) and width(grid[i + 1]) >= 2:
            return i
    for i in range(min(look, len(grid))):
        vals = [v for v in grid[i] if v]
        if len(vals) >= 2 and len(set(vals)) > 1 and all(width(grid[j]) <= len(vals) for j in range(i + 1, min(i + 3, len(grid)))):
            return i
    return next((i for i, r in enumerate(grid) if any(r)), 0)


def _sheet(ws) -> dict[str, Any]:
    nrows, ncols = min(ws.max_row, MAX_ROWS), min(ws.max_column, MAX_COLS)
    grid = [[_clean(c.value) for c in row] for row in ws.iter_rows(min_row=1, max_row=nrows, max_col=ncols)]
    flat = " ".join(v for r in grid for v in r if v)  # raw values before unmerging, for the naive baseline
    merged = []
    for rng in ws.merged_cells.ranges:  # unmerge everywhere, not just the top rows
        value = grid[rng.min_row - 1][rng.min_col - 1] if rng.min_row <= nrows and rng.min_col <= ncols else None
        merged.append(str(rng))
        for r in range(rng.min_row, min(rng.max_row, nrows) + 1):
            for c in range(rng.min_col, min(rng.max_col, ncols) + 1):
                grid[r - 1][c - 1] = value
    h = _header_row(grid)
    used = [c for c in range(ncols) if any(r[c] for r in grid[h:])]
    headers = {c: grid[h][c] or f"Column {get_column_letter(c + 1)}" for c in used}  # keyed by real position
    description = " ".join(dict.fromkeys(v for r in grid[:h] for v in r if v)) or None
    links = {(cell.row, cell.column): cell.hyperlink.target for row in ws.iter_rows(max_row=nrows, max_col=ncols)
             for cell in row if cell.hyperlink and cell.hyperlink.target}
    rows = []
    for r in range(h + 1, len(grid)):
        fields = {headers[c]: grid[r][c] for c in used if grid[r][c]}
        if fields:
            rows.append({"row": r + 1, "fields": fields,
                         "links": {headers[c - 1]: u for (rr, c), u in links.items() if rr == r + 1 and c - 1 in headers}})
    return {"name": ws.title, "hidden": ws.sheet_state != "visible", "header_row": h + 1,
            "headers": [headers[c] for c in used], "description": description, "merged": merged,
            "rows": rows, "preview": grid[:40], "flat": flat, "truncated": ws.max_row > MAX_ROWS or ws.max_column > MAX_COLS}


def extract(data) -> dict[str, Any]:
    wb = load_workbook(data, data_only=True)  # loaded once (it used to be reopened for every sheet)
    return {"kind": "xlsx", "sheets": [_sheet(ws) for ws in wb.worksheets]}
